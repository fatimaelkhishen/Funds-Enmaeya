#!/usr/bin/env python3
"""
run_all.py
----------
Runs every per-organization scraper in scraper/sites/ one after another
(each site keeps its own script, because every website is built differently),
then prints a summary table and saves it to scraper/run_summary.csv.

TEST FIRST (safe - writes to News/_test_output/, never touches your real
spreadsheets or the "already scraped" memory):

    python run_all.py --test                       # every site, 1 page, 2 articles each
    python run_all.py --test --only UNRWA Ministry_of_Labor

Real runs:

    python run_all.py                              # everything, full crawl
    python run_all.py --incremental                # weekly mode: stop each site once it reaches
                                                   # pages that were already scraped
    python run_all.py --only UNDP_Lebanon Ministry_of_Culture
    python run_all.py --skip National_News_Agency IOM
    python run_all.py --max-pages 20 --max-articles 50 --delay 2
    python run_all.py --list                       # just list the site scripts
    python run_all.py --incremental --daily-dir scrapes/2026-10-27
                                                   # also copy this run's NEW articles into
                                                   # scrapes/2026-10-27/<Org>.xlsx (used by the
                                                   # daily GitHub Actions run)

--only / --skip accept the file name with or without ".py", any letter case,
or just part of the name (e.g. --only ministry_of runs all ministries).
"""

import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import time
from datetime import datetime

SCRAPER_DIR = os.path.dirname(os.path.abspath(__file__))
SITES_DIR = os.path.join(SCRAPER_DIR, "sites")
NEWS_DIR = os.path.dirname(SCRAPER_DIR)
SUMMARY_CSV = os.path.join(SCRAPER_DIR, "run_summary.csv")


def output_file_of(script_path):
    """Read OUTPUT_FILE = '...' from a site script without importing it."""
    try:
        with open(script_path, encoding="utf-8") as f:
            m = re.search(r"""^OUTPUT_FILE\s*=\s*['"]([^'"]+)['"]""", f.read(), re.M)
        return m.group(1) if m else None
    except OSError:
        return None


def count_rows(xlsx_path):
    if not xlsx_path or not os.path.exists(xlsx_path):
        return 0
    try:
        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path, read_only=True)
        n = max((wb.active.max_row or 1) - 1, 0)
        wb.close()
        return n
    except Exception:
        return -1


def copy_new_rows(src_path, first_new_row, dest_path):
    """Copy rows first_new_row..end of src_path (plus its header row) into a new
    workbook at dest_path. Used by --daily-dir to keep one folder per run that
    holds only the articles found in that run."""
    from openpyxl import Workbook, load_workbook
    src = load_workbook(src_path, read_only=True)
    ws_src = src.active
    wb = Workbook()
    ws = wb.active
    ws.title = ws_src.title or "Articles"
    copied = 0
    for i, row in enumerate(ws_src.iter_rows(values_only=True), 1):
        if i == 1 or i >= first_new_row:
            ws.append(list(row))
            copied += i != 1
    src.close()
    os.makedirs(os.path.dirname(dest_path), exist_ok=True)
    wb.save(dest_path)
    return copied


def matches(script, names):
    strip = lambda x: x[:-3] if x.endswith(".py") else x
    s = strip(script.lower())
    for n in names:
        n = strip(n.lower())
        if s == n or n in s:
            return True
    return False


def main():
    p = argparse.ArgumentParser(description="Run every scraper in scraper/sites/ and print a summary.")
    p.add_argument("--test", action="store_true",
                   help="Quick safe test: 1 listing page and 2 articles per site, written to News/_test_output/.")
    p.add_argument("--dry-run", action="store_true",
                   help="Only check that each scraper finds article links and moves through the listing "
                        "pages: logs only, no articles opened, no Excel written (default 5 pages per site).")
    p.add_argument("--incremental", action="store_true",
                   help="Stop each site after 2 listing pages in a row with nothing new (good for weekly runs).")
    p.add_argument("--max-pages", type=int, default=None)
    p.add_argument("--max-articles", type=int, default=None)
    p.add_argument("--delay", type=float, default=None)
    p.add_argument("--only", nargs="*", help="Only run these sites (names or parts of names).")
    p.add_argument("--skip", nargs="*", help="Skip these sites.")
    p.add_argument("--timeout", type=int, default=None,
                   help="Max minutes per site before it is stopped (default 240, or 10 in --test).")
    p.add_argument("--list", action="store_true", help="List the site scripts and exit.")
    p.add_argument("--daily-dir", default=None,
                   help="Also save each site's NEW articles from this run into this folder "
                        "(relative to the repo root), e.g. scrapes/2026-10-27.")
    args = p.parse_args()

    scripts = sorted(f for f in os.listdir(SITES_DIR) if f.endswith(".py") and not f.startswith("_"))
    if args.only:
        scripts = [s for s in scripts if matches(s, args.only)]
    if args.skip:
        scripts = [s for s in scripts if not matches(s, args.skip)]

    if args.list:
        for s in scripts:
            print(f"{s:50s} -> {output_file_of(os.path.join(SITES_DIR, s))}")
        print(f"\n{len(scripts)} site script(s).")
        return
    if not scripts:
        print("No site scripts matched. Use --list to see the names.")
        sys.exit(1)

    env = dict(os.environ, SCRAPER_UNATTENDED="1", PYTHONIOENCODING="utf-8")
    extra = []
    max_pages, max_articles = args.max_pages, args.max_articles
    out_dir = NEWS_DIR
    if args.test:
        env["SCRAPER_TEST"] = "1"
        out_dir = os.path.join(NEWS_DIR, "_test_output")
        max_pages = 1 if max_pages is None else max_pages
        max_articles = 2 if max_articles is None else max_articles
        # Fresh, empty test state + output every time, so a test always really fetches.
        for d in (os.path.join(SCRAPER_DIR, "state_test"), out_dir):
            shutil.rmtree(d, ignore_errors=True)
        os.makedirs(out_dir, exist_ok=True)
    if args.dry_run:
        env["SCRAPER_DRY_RUN"] = "1"
        env.pop("SCRAPER_TEST", None)
        max_pages = 5 if max_pages is None else max_pages
        if args.delay is None:
            args.delay = 1.0
        shutil.rmtree(os.path.join(SCRAPER_DIR, "state_dry"), ignore_errors=True)
    if args.incremental and not args.dry_run:
        env["SCRAPER_STOP_AFTER_KNOWN_PAGES"] = "2"
    if max_pages is not None:
        extra += ["--max-pages", str(max_pages)]
    if max_articles is not None:
        extra += ["--max-articles", str(max_articles)]
    if args.delay is not None:
        extra += ["--delay", str(args.delay)]
    timeout_min = args.timeout if args.timeout is not None else (10 if (args.test or args.dry_run) else 240)

    state_dir = os.path.join(SCRAPER_DIR, "state_dry" if args.dry_run else ("state_test" if args.test else "state"))
    results = []
    started = time.time()
    mode = ("DRY RUN (pagination check - nothing is saved)" if args.dry_run
            else "TEST MODE - output goes to " + out_dir if args.test else "Running")
    print(f"{mode}: {len(scripts)} site(s), max pages {max_pages or 'default'}, "
          f"timeout {timeout_min} min each.")

    for i, script in enumerate(scripts, 1):
        script_path = os.path.join(SITES_DIR, script)
        out_file = output_file_of(script_path)
        out_path = os.path.join(out_dir, out_file) if out_file else None
        org_key = os.path.splitext(out_file)[0] if out_file else script[:-3]
        result_json = os.path.join(state_dir, "_last_run", f"{org_key}.json")
        before = count_rows(out_path)
        t0 = time.time()

        print(f"\n{'#' * 70}\n# [{i}/{len(scripts)}] {script}\n{'#' * 70}", flush=True)
        status, note = "ok", ""
        try:
            rc = subprocess.run([sys.executable, script_path] + extra, env=env,
                                stdin=subprocess.DEVNULL, timeout=timeout_min * 60).returncode
            if rc != 0:
                status, note = "CRASHED", f"exit code {rc} - see the error above / scrape_log"
        except subprocess.TimeoutExpired:
            status, note = "TIMEOUT", f"stopped after {timeout_min} min (articles saved so far are kept)"
        except KeyboardInterrupt:
            print("\nInterrupted by you - stopping here.")
            break

        after = count_rows(out_path)
        new_rows = after - before if before >= 0 and after >= 0 else 0
        if args.daily_dir and new_rows > 0 and not (args.test or args.dry_run):
            daily_path = os.path.join(NEWS_DIR, args.daily_dir, out_file)
            try:
                n = copy_new_rows(out_path, before + 2, daily_path)
                print(f"Saved {n} new article(s) to {daily_path}")
            except Exception as exc:
                print(f"(Could not write daily file {daily_path}: {exc})")
        info = {}
        if os.path.exists(result_json) and os.path.getmtime(result_json) >= t0 - 1:
            try:
                with open(result_json, encoding="utf-8") as f:
                    info = json.load(f)
            except Exception:
                info = {}
        if status == "ok":
            st = info.get("status", "ok")
            if st in ("blocked", "error", "skipped"):
                status, note = st.upper(), info.get("message", "")
            elif new_rows > 0:
                status = "OK"
            else:
                status = "NO NEW"
                if info.get("challenges"):
                    note = f"{info['challenges']} page(s) were bot-check/error pages"
                elif info and not info.get("links_found"):
                    note = "no article links found on the listing page"
                elif not args.test:
                    note = "nothing new since last run"
        if args.dry_run:
            dry = {}
            dry_json = os.path.join(state_dir, "_dry", f"{org_key}.json")
            if os.path.exists(dry_json):
                try:
                    with open(dry_json, encoding="utf-8") as f:
                        dry = json.load(f)
                except Exception:
                    dry = {}
            links = dry.get("links", 0)
            if status in ("ok", "OK", "NO NEW"):
                status = "OK" if links else "NO LINKS"
                note = ""
            if info.get("status") in ("blocked", "error"):
                status, note = info["status"].upper(), info.get("message", "")
            pages = info.get("pages", "")
            results.append({
                "site": script[:-3], "status": status, "pages": pages, "links": links,
                "minutes": round((time.time() - t0) / 60, 1),
                "first_link": (dry.get("first") or [""])[0], "last_link": (dry.get("last") or [""])[-1],
                "note": note,
            })
            continue
        results.append({
            "site": script[:-3], "status": status, "new_articles": new_rows, "total_rows": after,
            "minutes": round((time.time() - t0) / 60, 1), "note": note,
        })

    # ---- summary -----------------------------------------------------------
    if args.dry_run:
        print("\n" + "=" * 110)
        print(f"{'SITE':42s} {'STATUS':9s} {'PAGES':>5s} {'LINKS':>6s} {'MIN':>5s}  LAST LINK FOUND / NOTE")
        print("-" * 110)
        for r in results:
            print(f"{r['site'][:42]:42s} {r['status']:9s} {str(r['pages']):>5s} {r['links']:6d} "
                  f"{r['minutes']:5.1f}  {r['note'] or r['last_link'][:60]}")
        print("=" * 110)
        print(f"{sum(1 for r in results if r['status'] == 'OK')}/{len(results)} site(s) found article links. "
              "PAGES is blank for sites with their own custom pager - check scrape_log_dry.txt for those.\n"
              "Every link found per site: scraper/state_dry/_dry/<site>.json")
        path = os.path.join(SCRAPER_DIR, "pagination_check.csv")
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=list(results[0].keys()) if results else ["site"])
            w.writeheader()
            w.writerows(results)
        print(f"Summary saved to {path}\nFull log: scrape_log_dry.txt")
        return
    print("\n" + "=" * 100)
    print(f"{'SITE':45s} {'STATUS':9s} {'NEW':>5s} {'TOTAL':>6s} {'MIN':>5s}  NOTE")
    print("-" * 100)
    for r in results:
        print(f"{r['site'][:45]:45s} {r['status']:9s} {r['new_articles']:5d} {r['total_rows']:6d} "
              f"{r['minutes']:5.1f}  {r['note']}")
    print("=" * 100)
    good = sum(1 for r in results if r["status"] == "OK")
    print(f"{good}/{len(results)} site(s) saved new articles; "
          f"{sum(r['new_articles'] for r in results)} new article(s) in total; "
          f"{(time.time() - started) / 60:.1f} min.")

    summary_path = SUMMARY_CSV.replace(".csv", "_test.csv") if args.test else SUMMARY_CSV
    try:
        with open(summary_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=["run_at"] + list(results[0].keys()) if results else ["run_at"])
            w.writeheader()
            now = datetime.now().strftime("%Y-%m-%d %H:%M")
            for r in results:
                w.writerow(dict(r, run_at=now))
        print(f"Summary saved to {summary_path}")
        if args.daily_dir and not (args.test or args.dry_run):
            daily_dir = os.path.join(NEWS_DIR, args.daily_dir)
            os.makedirs(daily_dir, exist_ok=True)
            shutil.copy(summary_path, os.path.join(daily_dir, "run_summary.csv"))
    except Exception as exc:
        print(f"(Could not write summary file: {exc})")
    print("Full log: scrape_log" + ("_test" if args.test else "") + ".txt")


if __name__ == "__main__":
    main()
